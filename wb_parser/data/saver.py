import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from typing import List
from data.models import Product

class DataSaver:
    @staticmethod
    def save_to_csv(products: List[Product], filename: str = "output/csv/wildberries_products.csv"):
        """Сохранение в CSV"""
        if not products:
            print("Нет данных для сохранения")
            return
        
        fieldnames = ["Время", "Бренд", "Название", "Рейтинг", "Количество оценок", 
                     "Цена со скидкой", "Цена без скидки", "Скидка", "Ссылка на товар"]
        
        try:
            with open(filename, mode='w', newline='', encoding='utf-8') as file:
                writer = csv.DictWriter(file, fieldnames=fieldnames, delimiter='\t', quoting=csv.QUOTE_MINIMAL)
                writer.writeheader()
                
                for product in products:
                    writer.writerow({
                        "Время": product.timestamp,
                        "Бренд": product.brand,
                        "Название": product.name,
                        "Рейтинг": product.rate,
                        "Количество оценок": product.count_evolutions,
                        "Цена со скидкой": product.price,
                        "Цена без скидки": product.old_price,
                        "Скидка": product.discount,
                        "Ссылка на товар": product.link,
                    })
            
            print(f"\nДанные сохранены в {filename} (разделитель: табуляция)")
        except Exception as e:
            print(f"Ошибка при сохранении CSV: {e}")

    @staticmethod
    def save_to_excel(products: List[Product], filename: str = "output/excel/wildberries.xlsx"):
        """Сохранение в Excel"""
        if not products:
            print("Нет данных для сохранения в Excel")
            return
        
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Товары Wildberries"
            
            headers = [
                "Время", "Бренд", "Название", 
                "Рейтинг", "Количество оценок",
                "Цена со скидкой", "Цена без скидки", 
                "Скидка", "Ссылка на товар"
            ]
            
            # Настройка стилей
            header_font = Font(bold=True)
            align_center = Alignment(horizontal='center')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Заголовки
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = header_font
                cell.alignment = align_center
                cell.border = thin_border
            
            # Данные
            for row_num, product in enumerate(products, 2):
                ws.cell(row=row_num, column=1, value=product.timestamp).border = thin_border
                ws.cell(row=row_num, column=2, value=product.brand).border = thin_border
                ws.cell(row=row_num, column=3, value=product.name).border = thin_border
                ws.cell(row=row_num, column=4, value=product.rate).border = thin_border
                ws.cell(row=row_num, column=5, value=product.count_evolutions).border = thin_border
                ws.cell(row=row_num, column=6, value=product.price).border = thin_border
                ws.cell(row=row_num, column=7, value=product.old_price if product.old_price else "").border = thin_border
                ws.cell(row=row_num, column=8, value=product.discount if product.discount else "").border = thin_border
                ws.cell(row=row_num, column=9, value=product.link).border = thin_border
            
            # Ширина столбцов
            column_widths = {
                'A': 20, 'B': 20, 'C': 50, 
                'D': 10, 'E': 15, 'F': 15,
                'G': 15, 'H': 10, 'I': 80
            }
            
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
            
            wb.save(filename)
            print(f"\nДанные успешно сохранены в Excel файл: {filename}")
            
        except Exception as e:
            print(f"Ошибка при сохранении Excel: {e}")